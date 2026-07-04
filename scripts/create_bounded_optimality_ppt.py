from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs"
OUT_PATH = OUT_DIR / "demt_rebuttal_bounded_optimality_report.pptx"


COLORS = {
    "ink": RGBColor(31, 41, 55),
    "muted": RGBColor(107, 114, 128),
    "line": RGBColor(209, 213, 219),
    "light": RGBColor(243, 244, 246),
    "blue": RGBColor(37, 99, 235),
    "green": RGBColor(22, 163, 74),
    "amber": RGBColor(217, 119, 6),
    "red": RGBColor(220, 38, 38),
    "purple": RGBColor(124, 58, 237),
    "slate": RGBColor(15, 23, 42),
    "white": RGBColor(255, 255, 255),
}


def add_textbox(slide, x, y, w, h, text, size=16, color=None, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    if align is not None:
        p.alignment = align
    for paragraph in tf.paragraphs:
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color or COLORS["ink"]
    return shape


def add_title(slide, title, subtitle=None):
    add_textbox(slide, 0.55, 0.28, 12.2, 0.48, title, size=25, bold=True)
    if subtitle:
        add_textbox(slide, 0.57, 0.83, 12.1, 0.32, subtitle, size=11, color=COLORS["muted"])
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.20), Inches(12.2), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = COLORS["line"]
    line.line.fill.background()


def add_footer(slide):
    add_textbox(slide, 0.55, 7.13, 8.6, 0.2, "DEMT rebuttal bounded-optimality path", size=8, color=COLORS["muted"])


def add_bullets(slide, x, y, w, h, bullets, size=15, color=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = color or COLORS["ink"]
        p.space_after = Pt(5)
    return shape


def add_box(slide, x, y, w, h, title, body, fill, title_size=14, body_size=12):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.08)
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(title_size)
    p.font.bold = True
    p.font.color.rgb = COLORS["white"]
    p2 = tf.add_paragraph()
    p2.text = body
    p2.font.size = Pt(body_size)
    p2.font.color.rgb = COLORS["white"]
    return shape


def add_table(slide, x, y, w, h, rows, col_widths=None, size=10):
    table_shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    if col_widths:
        for idx, width in enumerate(col_widths):
            table.columns[idx].width = Inches(width)
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLORS["light"] if r_idx == 0 else COLORS["white"]
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.05)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(size)
                paragraph.font.bold = r_idx == 0
                paragraph.font.color.rgb = COLORS["ink"]
                paragraph.alignment = PP_ALIGN.CENTER if c_idx > 0 else PP_ALIGN.LEFT
    return table_shape


def add_bar_chart(slide, x, y, w, h, title, labels, values, color):
    data = CategoryChartData()
    data.categories = labels
    data.add_series(title, values)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
        data,
    ).chart
    chart.has_legend = False
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = 1
    chart.value_axis.tick_labels.font.size = Pt(9)
    chart.category_axis.tick_labels.font.size = Pt(9)
    chart.plots[0].has_data_labels = True
    chart.plots[0].data_labels.number_format = "0%"
    chart.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    chart.plots[0].data_labels.font.size = Pt(9)
    series = chart.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = color
    return chart


def add_percent_chart(slide, x, y, w, h, title, labels, values, color):
    add_textbox(slide, x, y - 0.28, w, 0.24, title, size=12, bold=True)
    return add_bar_chart(slide, x, y, w, h, title, labels, values, color)


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1. Title
    slide = prs.slides.add_slide(blank)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLORS["slate"]
    bg.line.fill.background()
    add_textbox(slide, 0.75, 1.15, 11.8, 0.75, "DEMT Rebuttal: Bounded Optimality Path", size=31, color=COLORS["white"], bold=True)
    add_textbox(slide, 0.78, 2.05, 11.5, 0.55, "Table 5 AR guidance parameters: defensible finite candidates, not global optima", size=20, color=RGBColor(191, 219, 254), bold=True)
    add_bullets(
        slide,
        0.82,
        3.05,
        11.5,
        1.8,
        [
            "Scope: spatial 25/5/3 cm funnel, 15 deg grasp-orientation tolerance, [0.5, 2.0] x v_ref pacing cue.",
            "Method: three-agent audit of Week1/Week2 evidence, human diagnostics, and experiment designs.",
            "Output: claim boundary, accepted rebuttal wording, and E05 fixed-ratio spatial funnel sweep plan.",
        ],
        size=17,
        color=COLORS["white"],
    )
    add_textbox(slide, 0.8, 6.52, 10.5, 0.28, "Generated from rebuttal_workflow/experiment_queue.md and updated handoffs", size=11, color=RGBColor(209, 213, 219))

    # 2. Reviewer pressure
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Reviewer pressure", "The question is parameter defensibility, not whether P/R/V structures matter")
    add_box(slide, 0.75, 1.55, 3.8, 1.35, "Reviewer attack", "Why are the submitted AR guidance constants close to optimal rather than arbitrary?", COLORS["red"])
    add_box(slide, 4.8, 1.55, 3.8, 1.35, "Weak answer", "Global search is impossible, so the values are implementation choices.", COLORS["amber"])
    add_box(slide, 8.85, 1.55, 3.8, 1.35, "Accepted answer", "Bounded finite-candidate optimality under the evaluated robot, task family, learner, protocol, and N=30 budget.", COLORS["green"])
    add_bullets(
        slide,
        0.85,
        3.35,
        11.8,
        2.55,
        [
            "Do not claim mathematical or global optimality.",
            "Do separate human compatibility from learner compatibility.",
            "Deployment success is primary; validation loss and structure metrics are diagnostics.",
            "If a new finite candidate dominates Table 5, the rebuttal claim must be downgraded.",
        ],
        size=17,
    )
    add_footer(slide)

    # 3. Evidence map
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Evidence map", "What can enter rebuttal now vs. what stays as limitation")
    rows = [
        ["Claim", "Status", "Evidence", "Boundary"],
        ["Bounded Table 5", "Usable", "EDSR, Table 9, finite sensitivity, human diagnostics", "Not global optimality"],
        ["Spatial 25/5/3", "Conservative", "EDSR + MPCV; funnel role separation", "No exact 25 cm proof"],
        ["Orientation 15 deg", "Strongest", "R00/R15/R30 + post p95 ~= 13-14 deg", "No R10/R20 bracket yet"],
        ["Temporal [0.5,2.0]", "Weak threshold", "MPSV + T00/T20/Twide + human coverage", "No clean ladder"],
    ]
    add_table(slide, 0.55, 1.55, 12.25, 2.35, rows, col_widths=[2.15, 1.55, 5.2, 3.35], size=10)
    add_box(slide, 0.75, 4.35, 5.7, 1.35, "Can say now", "Table 5 values are bounded finite candidates satisfying human-compatibility and learner-compatibility constraints.", COLORS["green"], body_size=13)
    add_box(slide, 6.9, 4.35, 5.7, 1.35, "Cannot say", "25 cm, 15 deg, and [0.5,2.0] x v_ref are exact or globally optimal thresholds.", COLORS["red"], body_size=13)
    add_footer(slide)

    # 4. Paper-level anchor
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Strongest existing evidence", "The paper-level real-robot results anchor the bounded claim")
    add_percent_chart(slide, 0.75, 1.85, 5.4, 3.05, "Deployment success: target users", ["P1 pre", "P6 post", "P2 pre", "P7 post"], [0.14, 0.75, 0.22, 0.91], COLORS["blue"])
    rows = [
        ["Metric", "Pre", "Post DEMT", "Interpretation"],
        ["MPCV spatial", "P1 0.0046; P2 0.0083", "P6 0.0007; P7 0.0005", "Full-trajectory path tightens"],
        ["MCOD orientation", "P1 0.3342; P2 0.6582", "P6 0.0749; P7 0.1668", "Contact orientation tightens"],
        ["MPSV temporal", "P1 0.0026; P2 0.0056", "P6 0.0013; P7 0.0010", "Pacing synchronizes"],
    ]
    add_table(slide, 6.55, 1.85, 6.1, 2.55, rows, col_widths=[1.35, 1.55, 1.65, 1.55], size=8.5)
    add_bullets(
        slide,
        6.65,
        4.78,
        5.8,
        1.3,
        [
            "This supports learner-compatible structure after guidance removal.",
            "It does not isolate exact Table 5 thresholds.",
        ],
        size=14,
    )
    add_footer(slide)

    # 5. Spatial audit
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Spatial audit", "P00/P01/P10/P11 is valid history, but weak for 25 cm optimality")
    add_percent_chart(slide, 0.75, 1.70, 5.2, 2.95, "Week1 corrected spatial rollout", ["P00", "P01", "P10", "P11"], [0.8, 0.6, 0.7, 0.6], COLORS["amber"])
    rows = [
        ["Condition", "Corridor", "Pre-grasp", "Result"],
        ["P00", "10 cm", "2 cm", "8/10"],
        ["P01", "10 cm", "6 cm", "6/10"],
        ["P10", "25 cm", "2 cm", "7/10; worse valid loss"],
        ["P11", "25 cm", "6 cm", "6/10"],
    ]
    add_table(slide, 6.35, 1.55, 6.25, 2.25, rows, col_widths=[1.25, 1.35, 1.45, 2.2], size=10)
    add_bullets(
        slide,
        6.48,
        4.25,
        5.95,
        1.55,
        [
            "Independent factors do not test the coupled Table 5 funnel.",
            "Rollout is mixed; do not use this as monotonic spatial proof.",
            "Use paper EDSR/MPCV plus funnel role separation for current rebuttal.",
        ],
        size=14,
    )
    add_footer(slide)

    # 6. E00 human compatibility
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Human compatibility bounds", "E00 is useful support, not optimality evidence")
    add_percent_chart(slide, 0.65, 1.68, 5.8, 2.75, "Pre-training spatial coverage", ["P1 10", "P1 25", "P2 10", "P2 25", "P2 30"], [0.65, 1.0, 0.325, 0.908, 0.983], COLORS["purple"])
    rows = [
        ["Human diagnostic", "Key number", "Use"],
        ["P1 spatial p95", "18.28 cm", "25 cm is permissive"],
        ["P2 spatial p95", "26.70 cm", "25 cm covers most, not all"],
        ["P6/P7 post spatial p95", "9.49 / 6.67 cm", "Training tightens inside envelope"],
        ["P6/P7 orientation p95", "13-14 deg", "Supports 15 deg compatibility"],
        ["Temporal [0.5,2.0]", "86.7-100% duration; 93.3-100% speed", "Broad pacing scaffold"],
    ]
    add_table(slide, 6.75, 1.55, 5.85, 3.05, rows, col_widths=[2.15, 1.5, 2.2], size=8.8)
    add_box(slide, 1.0, 5.25, 11.4, 0.85, "Important limitation", "These are observational compatibility bounds. They cannot prove learner or Pareto optimality without matched policy training/evaluation.", COLORS["amber"], body_size=12)
    add_footer(slide)

    # 7. Orientation
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Orientation is the strongest parameter story", "15 deg has finite-candidate learner and human support")
    add_percent_chart(slide, 0.75, 1.70, 5.35, 2.95, "Orientation rollout", ["R00", "R15", "R30"], [0.7, 0.6, 0.4], COLORS["green"])
    add_percent_chart(slide, 6.85, 1.70, 5.35, 2.95, "Post-training 15 deg coverage", ["P6 grasp", "P6 close", "P7 grasp", "P7 close"], [1.0, 1.0, 0.992, 0.992], COLORS["blue"])
    add_bullets(
        slide,
        0.95,
        5.05,
        11.4,
        1.0,
        [
            "Allowed claim: 15 deg is human-achievable and more learner-compatible than 30 deg in the tested finite set.",
            "Limitation: no R10/R20 local bracket or exact target-frame human alignment yet.",
        ],
        size=14,
    )
    add_footer(slide)

    # 8. Temporal
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Temporal evidence needs conservative wording", "T00/T20/Twide supports pacing structure, not exact [0.5, 2.0] optimality")
    add_percent_chart(slide, 0.75, 1.75, 5.4, 2.95, "Temporal rollout", ["T00", "T20", "Twide"], [1.0, 0.8, 0.7], COLORS["amber"])
    add_box(slide, 6.65, 1.72, 5.7, 1.3, "What it supports", "Temporal consistency is learner-relevant; broad phase variation is riskier under N=30.", COLORS["green"])
    add_box(slide, 6.65, 3.32, 5.7, 1.3, "What it does not prove", "The submitted [0.5, 2.0] x v_ref interval is not exactly optimized.", COLORS["red"])
    add_bullets(
        slide,
        6.75,
        5.05,
        5.5,
        1.0,
        [
            "Backup experiment: clean T00/T10/T20/T30/T40/T50 ladder.",
            "Human ratios are to group median, not per-user v_ref.",
        ],
        size=13,
    )
    add_footer(slide)

    # 9. Agent3 criteria
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Agent3 acceptance criteria", "What a stronger bounded-optimality claim must satisfy")
    rows = [
        ["Criterion", "Requirement"],
        ["Finite set", "Table 5 value plus tighter and looser alternatives"],
        ["Controls", "Same robot, task family, learner, N=30, checkpoint rule, eval starts"],
        ["Primary metric", "Deployment success; validation and structure metrics are secondary"],
        ["Human vs learner", "Evaluate human compatibility separately from policy learnability"],
        ["Pareto wording", "Only if no candidate dominates on both axes"],
        ["Claim discipline", "If another candidate wins, downgrade the rebuttal"],
    ]
    add_table(slide, 1.0, 1.55, 11.35, 3.8, rows, col_widths=[2.5, 8.85], size=12)
    add_footer(slide)

    # 10. E05 plan
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Highest-value next experiment: E05", "Fixed-ratio spatial funnel sweep")
    rows = [
        ["Candidate", "True 25/5/3 family", "Proxy corridor:pre-grasp family"],
        ["Tighter", "20 / 4 / 2.4 cm", "20 : 4.8 cm"],
        ["Table 5 scale", "25 / 5 / 3.0 cm", "25 : 6.0 cm"],
        ["Looser", "30 / 6 / 3.6 cm", "30 : 7.2 cm"],
        ["Optional", "15 / 3 / 1.8; 35 / 7 / 4.2", "15 : 3.6; 35 : 8.4"],
    ]
    add_table(slide, 0.8, 1.55, 11.75, 2.25, rows, col_widths=[1.75, 4.9, 5.1], size=10.5)
    add_bullets(
        slide,
        1.0,
        4.25,
        11.25,
        1.55,
        [
            "Fix R and V; N=30 per condition; same learner, training budget, checkpoint rule, and paired rollout starts.",
            "Positive result: S25 is top or tied and no alternative dominates on human compatibility plus deployment.",
            "Negative result: keep spatial as a useful funnel/structure claim and list threshold optimization as future work.",
        ],
        size=14,
    )
    add_footer(slide)

    # 11. Queue
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Experiment queue", "Use compute only where it can change rebuttal wording")
    rows = [
        ["ID", "Experiment", "Relevance", "Evidence if positive"],
        ["E00", "Human percentile bound audit", "Done / immediate", "Human compatibility support"],
        ["E05", "Fixed-ratio spatial funnel sweep", "Immediate if compute allows", "Strong spatial bounded-candidate claim"],
        ["E06", "Clean temporal ladder", "Backup", "Temporal sensitivity, not exact v_ref proof"],
        ["E07", "R10/R20 orientation bracket", "Backup / future", "Stronger local 15 deg claim"],
        ["E08", "Human-derived P/R/V sampling", "Future paper", "Data-calibrated DEMT extension"],
    ]
    add_table(slide, 0.75, 1.55, 11.9, 3.15, rows, col_widths=[0.7, 3.2, 2.45, 5.55], size=10)
    add_box(slide, 1.1, 5.35, 11.1, 0.82, "Decision", "Do not start broad P/R/V sampling for rebuttal unless E05 is rejected or impossible and Agent3 demands more.", COLORS["purple"], body_size=12)
    add_footer(slide)

    # 12. Final wording
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Rebuttal wording", "Concise claim boundary")
    add_box(
        slide,
        0.85,
        1.55,
        11.65,
        1.55,
        "Allowed",
        "The Table 5 values instantiate DEMT-selected P/R/V structures and are defended as bounded finite candidates under human-compatibility and learner-compatibility constraints for this robot, task family, learner, and N=30 budget.",
        COLORS["green"],
        body_size=12,
    )
    add_box(
        slide,
        0.85,
        3.55,
        11.65,
        1.3,
        "Forbidden",
        "The Table 5 thresholds are globally optimal, exactly optimized, or fully justified by the current follow-up experiments.",
        COLORS["red"],
        body_size=12,
    )
    add_bullets(
        slide,
        1.05,
        5.35,
        11.2,
        0.85,
        [
            "Spatial: role-separated 25/5/3 funnel, not exact 25 cm proof.",
            "Orientation: strongest current finite candidate.",
            "Temporal: broad scaffold and limitation until clean ladder.",
        ],
        size=14,
    )
    add_footer(slide)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    print(build_deck())
