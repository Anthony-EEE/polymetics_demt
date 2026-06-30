from pathlib import Path
import json

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs"
OUT_PATH = OUT_DIR / "demt_rebuttal_mvp_report.pptx"

DATA_ROOT = Path("/scratch/prj/eng_demt_robot_learning/polymetics_demt/dataset")
SPATIAL_SUMMARY = DATA_ROOT / "abla1_full_arcenter/policy_rollouts_seed628_latest/summary_seed628_n10.json"
ORIENTATION_SUMMARY = DATA_ROOT / "orn_mvp_full/policy_rollouts_seed628/summary_seed628_n10.json"
TEMPORAL_SUMMARY = DATA_ROOT / "time_mvp_full/policy_rollouts_seed628_latest/summary_seed628_n10.json"

LOSS_PNG = DATA_ROOT / "abla1_full_arcenter/loss.png"
SPATIAL_P00_PNG = DATA_ROOT / "abla1_full_arcenter/trajectory_plots/abla1_P00_ee_xyz_trajectories.png"
SPATIAL_P10_PNG = DATA_ROOT / "abla1_full_arcenter/trajectory_plots/abla1_P10_ee_xyz_trajectories.png"
ORN_R30_PNG = DATA_ROOT / "orn_mvp_full/orientation_plots/orn_mvp_R30_ee_xyz_orientation.png"
TIME_SUMMARY_PNG = DATA_ROOT / "time_mvp_full/temporal_plots/time_mvp_temporal_summary.png"


COLORS = {
    "ink": RGBColor(31, 41, 55),
    "muted": RGBColor(107, 114, 128),
    "light": RGBColor(243, 244, 246),
    "line": RGBColor(209, 213, 219),
    "blue": RGBColor(37, 99, 235),
    "green": RGBColor(22, 163, 74),
    "amber": RGBColor(217, 119, 6),
    "red": RGBColor(220, 38, 38),
    "purple": RGBColor(124, 58, 237),
    "white": RGBColor(255, 255, 255),
}


def load_rates(path):
    with path.open() as f:
        data = json.load(f)
    rows = []
    for item in data["summaries"]:
        rows.append(
            {
                "condition": item["condition"],
                "success": item["success_count"],
                "total": item["num_rollouts"],
                "rate": item["success_rate"],
                "config": item.get("condition_config", {}),
                "checkpoint": Path(item["checkpoint"]).name if "checkpoint" in item else "",
            }
        )
    return rows


def set_text_frame(tf, font_size=18, color=None, bold=False):
    for paragraph in tf.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(font_size)
            run.font.bold = bold
            if color:
                run.font.color.rgb = color


def add_textbox(slide, x, y, w, h, text, size=18, color=None, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    if align is not None:
        p.alignment = align
    set_text_frame(tf, size, color or COLORS["ink"], bold)
    return shape


def add_title(slide, title, subtitle=None):
    add_textbox(slide, 0.55, 0.28, 12.2, 0.55, title, size=26, bold=True)
    if subtitle:
        add_textbox(slide, 0.57, 0.86, 12.0, 0.35, subtitle, size=12, color=COLORS["muted"])
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.23), Inches(12.2), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = COLORS["line"]
    line.line.fill.background()


def add_footer(slide, text="DEMT rebuttal MVP report"):
    add_textbox(slide, 0.55, 7.12, 8.0, 0.2, text, size=8, color=COLORS["muted"])


def add_bullets(slide, x, y, w, h, bullets, size=16, color=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = color or COLORS["ink"]
        p.space_after = Pt(6)
    return shape


def add_callout(slide, x, y, w, h, title, body, fill, title_size=14, body_size=13):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = fill
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.14)
    tf.margin_right = Inches(0.14)
    tf.margin_top = Inches(0.08)
    p = tf.paragraphs[0]
    p.text = title
    p.font.bold = True
    p.font.size = Pt(title_size)
    p.font.color.rgb = COLORS["white"]
    p2 = tf.add_paragraph()
    p2.text = body
    p2.font.size = Pt(body_size)
    p2.font.color.rgb = COLORS["white"]
    return box


def add_table(slide, x, y, w, h, rows, col_widths=None, font_size=11, header_fill=None):
    table_shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    if col_widths:
        for idx, width in enumerate(col_widths):
            table.columns[idx].width = Inches(width)
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = text
            fill = COLORS["light"] if r_idx == 0 else COLORS["white"]
            if r_idx == 0 and header_fill:
                fill = header_fill
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.05)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(font_size)
                paragraph.font.color.rgb = COLORS["ink"] if r_idx != 0 or not header_fill else COLORS["white"]
                paragraph.font.bold = r_idx == 0
                paragraph.alignment = PP_ALIGN.CENTER if c_idx > 0 else PP_ALIGN.LEFT
    return table_shape


def add_picture_fit(slide, image_path, x, y, w, h):
    image_path = Path(image_path)
    with Image.open(image_path) as im:
        iw, ih = im.size
    box_ratio = w / h
    img_ratio = iw / ih
    if img_ratio >= box_ratio:
        final_w = w
        final_h = w / img_ratio
        final_x = x
        final_y = y + (h - final_h) / 2
    else:
        final_h = h
        final_w = h * img_ratio
        final_x = x + (w - final_w) / 2
        final_y = y
    return slide.shapes.add_picture(str(image_path), Inches(final_x), Inches(final_y), width=Inches(final_w), height=Inches(final_h))


def add_bar_chart(slide, x, y, w, h, title, labels, values, color):
    chart_data = CategoryChartData()
    chart_data.categories = labels
    chart_data.add_series(title, values)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
        chart_data,
    ).chart
    chart.has_legend = False
    chart.value_axis.maximum_scale = 1.0
    chart.value_axis.minimum_scale = 0.0
    chart.value_axis.tick_labels.font.size = Pt(9)
    chart.category_axis.tick_labels.font.size = Pt(10)
    chart.plots[0].has_data_labels = True
    data_labels = chart.plots[0].data_labels
    data_labels.number_format = "0%"
    data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    data_labels.font.size = Pt(10)
    series = chart.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = color
    return chart


def section_label(slide, x, y, text, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(1.6), Inches(0.28))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLORS["white"]


def build_deck():
    spatial = load_rates(SPATIAL_SUMMARY)
    orientation = load_rates(ORIENTATION_SUMMARY)
    temporal = load_rates(TEMPORAL_SUMMARY)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1
    slide = prs.slides.add_slide(blank)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(17, 24, 39)
    bg.line.fill.background()
    add_textbox(slide, 0.8, 1.25, 11.6, 0.9, "DEMT Rebuttal MVP Results", size=34, color=COLORS["white"], bold=True)
    add_textbox(slide, 0.82, 2.14, 10.8, 0.55, "Why these AR guidance parameters?", size=24, color=RGBColor(191, 219, 254), bold=True)
    add_bullets(
        slide,
        0.86,
        3.05,
        11.2,
        1.7,
        [
            "Question: are corridor width, grasp-orientation tolerance, and speed-bar thresholds arbitrary?",
            "New evidence: three controlled MVP ablations for spatial, orientation, and temporal structure.",
            "Main claim: the guidance structures are sensitivity-supported; exact numeric optimality is not claimed.",
        ],
        size=18,
        color=COLORS["white"],
    )
    add_textbox(slide, 0.84, 6.55, 10.5, 0.3, "Prepared for supervisor discussion", size=12, color=RGBColor(209, 213, 219))

    # 2
    slide = prs.slides.add_slide(blank)
    add_title(slide, "The rebuttal question", "What the reviewer / supervisor may ask")
    add_callout(
        slide,
        0.7,
        1.65,
        5.7,
        1.1,
        "Concern",
        "Why are the AR guidance parameters set to these values?",
        COLORS["red"],
        title_size=16,
        body_size=17,
    )
    add_callout(
        slide,
        6.85,
        1.65,
        5.7,
        1.1,
        "Stronger version",
        "Can we prove 25cm / 15deg / [0.5,2.0] are optimal?",
        COLORS["amber"],
        title_size=16,
        body_size=17,
    )
    add_bullets(
        slide,
        0.8,
        3.25,
        11.6,
        2.0,
        [
            "The original paper supports the structures: spatial path, contact orientation, and temporal pacing.",
            "It reports concrete AR implementation parameters, but does not run per-parameter threshold sweeps.",
            "The MVPs therefore should be framed as sensitivity evidence, not global parameter optimization.",
        ],
        size=18,
    )
    add_callout(
        slide,
        1.1,
        5.75,
        11.1,
        0.55,
        "Meeting target",
        "Decide whether this evidence is enough for rebuttal, or whether one final stress rollout is needed.",
        COLORS["blue"],
        title_size=13,
        body_size=13,
    )
    add_footer(slide)

    # 3
    slide = prs.slides.add_slide(blank)
    add_title(slide, "What the paper already establishes", "DEMT evidence chain in the submitted PDF")
    rows = [
        ["Step", "Role in the argument", "Evidence in paper"],
        ["1. Structure space", "Learner-relevant variations are spatial, contact, temporal.", "DEMT defines P x R x V."],
        ["2. Deployment selection", "Task success is not enough under fixed learner and budget.", "Simulation variants reduce deployment success."],
        ["3. Human guidance", "Selected structures become AR training targets.", "Corridor, grasp cue, speed bar."],
        ["4. Real study", "Guidance changes unguided teaching behavior.", "Target group outperforms control and ARCap baseline."],
    ]
    add_table(slide, 0.65, 1.55, 12.0, 2.55, rows, col_widths=[1.7, 5.2, 5.1], font_size=12, header_fill=COLORS["blue"])
    add_callout(
        slide,
        0.95,
        4.55,
        5.5,
        1.35,
        "Supported",
        "The dimensions are learner-relevant and human-trainable.",
        COLORS["green"],
        title_size=15,
        body_size=16,
    )
    add_callout(
        slide,
        6.9,
        4.55,
        5.5,
        1.35,
        "Not fully shown",
        "The exact numeric thresholds are not proven globally optimal.",
        COLORS["amber"],
        title_size=15,
        body_size=16,
    )
    add_footer(slide)

    # 4
    slide = prs.slides.add_slide(blank)
    add_title(slide, "AR guidance parameters in the PDF", "What Table 5 says and how we should interpret it")
    rows = [
        ["Guidance", "Reported parameter", "Interpretation"],
        ["Spatial corridor", "r_start=25cm, r_middle=5cm, r_final=3cm", "Funnel-like path constraint, not a single sampled radius."],
        ["Grasp orientation", "15 deg tolerance", "Allowed contact-orientation deviation around target grasp family."],
        ["Speed bar", "[0.5, 2.0] x v_ref", "Relative pacing cue normalized by reference speed."],
    ]
    add_table(slide, 0.75, 1.6, 11.8, 2.25, rows, col_widths=[2.2, 4.2, 5.4], font_size=12, header_fill=COLORS["purple"])
    add_bullets(
        slide,
        0.95,
        4.35,
        11.3,
        1.6,
        [
            "Rebuttal should avoid claiming that each value is uniquely optimal.",
            "Safer claim: the parameters implement DEMT-selected structures; the MVPs support the need for these constraints.",
            "This separates structure-level evidence from exact threshold optimality.",
        ],
        size=17,
    )
    add_footer(slide)

    # 5
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Three MVPs: what they answer", "Each MVP relaxes one learner-relevant structure")
    rows = [
        ["MVP", "Relaxed factor", "Conditions", "Question answered"],
        ["Spatial", "Corridor entry / pre-grasp spread", "P00, P01, P10, P11", "Does spatial corridor variability hurt learning?"],
        ["Orientation", "Full-trajectory grasp orientation", "R00, R15, R30", "Is 15deg a reasonable tolerance boundary?"],
        ["Temporal", "Phase-duration variation", "T00, T20, Twide", "Does broad timing variation hurt learning?"],
    ]
    add_table(slide, 0.5, 1.45, 12.3, 2.35, rows, col_widths=[1.5, 3.4, 2.6, 4.8], font_size=11, header_fill=COLORS["blue"])
    add_callout(
        slide,
        0.8,
        4.35,
        5.75,
        1.0,
        "Can support",
        "Relaxing these structures changes learner fit and/or rollout success.",
        COLORS["green"],
    )
    add_callout(
        slide,
        6.8,
        4.35,
        5.75,
        1.0,
        "Cannot prove",
        "The exact reported values are globally optimal across tasks, users, or robots.",
        COLORS["red"],
    )
    add_footer(slide)

    # 6
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Spatial MVP result (corrected)", "AR-center, no-random collection with latest-checkpoint rollout")
    labels = [r["condition"] for r in spatial]
    values = [r["rate"] for r in spatial]
    add_bar_chart(slide, 0.65, 1.45, 5.4, 3.35, "Rollout success", labels, values, COLORS["blue"])
    rows = [["Condition", "Corridor radius", "Pre-grasp radius", "Success"]]
    for r in spatial:
        cfg = r["config"]
        rows.append([
            r["condition"],
            f"{cfg['corridor_start_radius']:.2f} m",
            f"{cfg['pre_grasp_radius']:.2f} m",
            f"{r['success']}/{r['total']}",
        ])
    add_table(slide, 6.45, 1.55, 6.25, 2.15, rows, col_widths=[1.2, 1.75, 1.75, 1.2], font_size=11, header_fill=COLORS["blue"])
    add_bullets(
        slide,
        6.6,
        4.05,
        5.8,
        1.5,
        [
            "Corrected rollout: P00/P01/P10/P11 = 8/10, 6/10, 7/10, 6/10.",
            "P10 has the worst validation loss; rollout differences are mixed.",
            "Use as structure-level sensitivity evidence; do not cite the old wrong-center run.",
        ],
        size=15,
    )
    add_footer(slide)

    # 7
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Spatial diagnostics (corrected)", "Loss curves and trajectory checks for AR-center, no-random data")
    add_picture_fit(slide, LOSS_PNG, 0.65, 1.45, 5.9, 2.25)
    add_picture_fit(slide, SPATIAL_P00_PNG, 6.85, 1.38, 5.7, 1.95)
    add_picture_fit(slide, SPATIAL_P10_PNG, 6.85, 3.55, 5.7, 1.95)
    add_bullets(
        slide,
        0.85,
        4.05,
        5.4,
        1.4,
        [
            "Checker passed for 30 demos/condition, no random_start phase, center [0.3, 0.0, 0.5].",
            "Best valid loss: P00/P01/P10/P11 = 0.036, 0.029, 0.099, 0.029.",
            "Trajectory plots confirm the corrected corridor_start geometry and wider P10/P11 spread.",
        ],
        size=14,
    )
    add_footer(slide)

    # 8
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Orientation MVP result", "15deg is a plausible tolerance; 30deg degrades more")
    labels = [r["condition"] for r in orientation]
    values = [r["rate"] for r in orientation]
    add_bar_chart(slide, 0.65, 1.45, 5.2, 3.25, "Rollout success", labels, values, COLORS["purple"])
    rows = [["Condition", "Orientation range", "Success"]]
    for r in orientation:
        deg = r["config"]["orientation_max_degrees"]
        rows.append([r["condition"], f"+/- {deg:.0f} deg", f"{r['success']}/{r['total']}"])
    add_table(slide, 6.25, 1.55, 3.95, 1.75, rows, col_widths=[1.25, 1.65, 1.05], font_size=11, header_fill=COLORS["purple"])
    add_picture_fit(slide, ORN_R30_PNG, 6.3, 3.6, 5.85, 2.15)
    add_bullets(
        slide,
        0.9,
        4.95,
        4.9,
        1.0,
        [
            "R00/R15/R30 success: 7/10, 6/10, 4/10.",
            "Training validation loss is also monotonic.",
            "Best rebuttal use: orientation tolerance is sensitivity-checked; 15deg remains learnable while 30deg degrades.",
        ],
        size=14,
    )
    add_footer(slide)

    # 9
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Temporal MVP result", "Broad phase-duration variation reduces latest-checkpoint rollout success")
    labels = [r["condition"] for r in temporal]
    values = [r["rate"] for r in temporal]
    add_bar_chart(slide, 0.65, 1.45, 5.2, 3.25, "Rollout success", labels, values, COLORS["amber"])
    rows = [["Condition", "Duration multipliers", "Checkpoint", "Success"]]
    for r in temporal:
        lo, hi = r["config"]["duration_multiplier_range"]
        rows.append([r["condition"], f"[{lo:.1f}, {hi:.1f}]", r["checkpoint"], f"{r['success']}/{r['total']}"])
    add_table(slide, 6.15, 1.5, 6.55, 1.85, rows, col_widths=[1.15, 1.8, 2.3, 1.0], font_size=10, header_fill=COLORS["amber"])
    add_picture_fit(slide, TIME_SUMMARY_PNG, 6.15, 3.65, 6.55, 2.05)
    add_bullets(
        slide,
        0.88,
        4.85,
        4.9,
        1.0,
        [
            "Latest fixed-start: T00/T20/Twide = 10/10, 8/10, 7/10.",
            "Spatial path and orientation were fixed to isolate timing.",
            "This supports speed/phase consistency as learner-relevant.",
        ],
        size=14,
    )
    add_footer(slide)

    # 10
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Cross-MVP synthesis", "The three guidance dimensions are learner-relevant")
    rows = [
        ["Dimension", "Best controlled condition", "Relaxed condition", "Observed pattern", "Rebuttal use"],
        ["Spatial", "Small corridor entry", "Wide corridor entry", "P10 val loss rises; rollout mixed", "Corridor constraints matter"],
        ["Orientation", "0deg / 15deg", "30deg", "Monotonic rollout/loss degradation", "15deg tolerance is reasonable"],
        ["Temporal", "Fixed timing", "Broad [0.5,2.0]", "Latest rollout decreases 1.0 -> 0.7", "Speed-bar consistency matters"],
    ]
    add_table(slide, 0.45, 1.45, 12.45, 2.55, rows, col_widths=[1.45, 2.25, 2.1, 3.25, 3.4], font_size=10, header_fill=COLORS["green"])
    add_callout(
        slide,
        0.9,
        4.45,
        5.75,
        1.0,
        "Strong claim",
        "The AR parameters implement structures that affect learnability under the fixed learner and N=30 budget.",
        COLORS["green"],
    )
    add_callout(
        slide,
        6.9,
        4.45,
        5.75,
        1.0,
        "Avoid",
        "Do not claim these exact numeric thresholds are globally optimal.",
        COLORS["red"],
    )
    add_footer(slide)

    # 11
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Suggested rebuttal wording", "Careful claim that matches the evidence")
    quote = (
        "The exact AR values are task-specific implementation choices for this robot, learner, "
        "and manipulation setting. Our claim is not that these numbers are globally optimal. "
        "Rather, they implement the spatial, contact-orientation, and temporal structures selected "
        "by DEMT. To address the concern, we added controlled sensitivity ablations showing that "
        "relaxing these structures degrades learner fit and/or rollout success under the same "
        "fixed learner and N=30 budget."
    )
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.55), Inches(11.7), Inches(2.25))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS["light"]
    shape.line.color.rgb = COLORS["line"]
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.15)
    p = tf.paragraphs[0]
    p.text = quote
    p.font.size = Pt(20)
    p.font.color.rgb = COLORS["ink"]
    add_bullets(
        slide,
        1.0,
        4.25,
        11.0,
        1.35,
        [
            "Then report the three MVP trends in one compact table.",
            "Explicitly state that exact threshold optimality remains a task-specific design limitation.",
            "This gives the reviewer a concrete answer without overclaiming.",
        ],
        size=16,
    )
    add_footer(slide)

    # 12
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Decision for tomorrow", "What to ask the supervisor")
    rows = [
        ["Option", "Pros", "Cons", "Recommendation"],
        ["Use MVPs as rebuttal evidence", "Fast; directly addresses concern; no new training", "Not a proof of optimality", "Yes"],
        ["Run final stress rollout", "Cleaner deployment claim; tests robustness", "Costs time; may complicate story", "Only if reviewer wording demands it"],
        ["Run full parameter optimization", "Strongest exact-number proof", "Too large for rebuttal timeline", "No"],
    ]
    add_table(slide, 0.55, 1.5, 12.2, 2.35, rows, col_widths=[2.3, 3.5, 3.2, 3.2], font_size=11, header_fill=COLORS["blue"])
    add_callout(
        slide,
        0.85,
        4.45,
        11.6,
        0.95,
        "Recommended position",
        "Use the MVPs to support parameter rationale, and phrase exact values as task-specific thresholds rather than optimized constants.",
        COLORS["blue"],
    )
    add_bullets(
        slide,
        1.0,
        5.75,
        11.2,
        0.8,
        [
            "If the supervisor wants one more check: run a non-overwriting stress rollout for orientation/temporal only.",
            "Otherwise, spend time turning the MVP table into rebuttal text and appendix material.",
        ],
        size=15,
    )
    add_footer(slide)

    # 13
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Backup: exact result files", "Useful for checking numbers during the meeting")
    add_bullets(
        slide,
        0.75,
        1.45,
        11.9,
        4.7,
        [
            f"Spatial summary: {SPATIAL_SUMMARY}",
            f"Orientation summary: {ORIENTATION_SUMMARY}",
            f"Temporal latest summary: {TEMPORAL_SUMMARY}",
            f"Spatial loss: {LOSS_PNG}",
            f"Orientation plot example: {ORN_R30_PNG}",
            f"Temporal summary plot: {TIME_SUMMARY_PNG}",
        ],
        size=13,
    )
    add_footer(slide)

    OUT_DIR.mkdir(exist_ok=True)
    prs.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    path = build_deck()
    print(path)
