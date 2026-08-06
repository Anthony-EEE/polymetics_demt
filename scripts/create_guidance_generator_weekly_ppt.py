#!/usr/bin/env python3
"""Create the 2026-07-28 T-RO weekly meeting deck.

The deck is intentionally generated from the frozen Stage 2c summary values.
It uses only repository-local experiment evidence and code-native visuals.
"""

from __future__ import annotations

import argparse
import math
import os
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import LinearSegmentedColormap


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = (
    REPO_ROOT
    / "reports"
    / "meetings"
    / "2026-07-28-guidance-generator-weekly.pptx"
)
ASSET_DIR = (
    REPO_ROOT
    / "reports"
    / "meetings"
    / "assets"
    / "guidance_generator_weekly"
)

# Frozen Stage 2c values from:
# experiments/tro_stage2_position_composition/analysis/summary.md
STARTS = [15, 25, 35]
APPROACHES = [3.6, 6.0, 8.4]
COMBINED = np.array(
    [
        [0.4944, 0.3720, 0.2840],
        [0.4496, 0.4896, 0.2408],
        [0.2872, 0.3336, 0.2152],
    ]
)
INNER = np.array(
    [
        [0.6512, 0.4720, 0.3872],
        [0.5184, 0.5456, 0.2528],
        [0.3056, 0.3056, 0.2256],
    ]
)
OUTER = np.array(
    [
        [0.3376, 0.2720, 0.1808],
        [0.3808, 0.4336, 0.2288],
        [0.2688, 0.3616, 0.2048],
    ]
)
ID_RATE = np.array(
    [
        [0.7584, 0.5776, 0.4176],
        [0.4784, 0.4784, 0.2736],
        [0.2800, 0.3200, 0.1792],
    ]
)
OOD_RATE = np.array(
    [
        [0.2688, 0.2208, 0.1488],
        [0.3168, 0.3600, 0.1776],
        [0.2512, 0.3280, 0.1632],
    ]
)

# Deck palette
NAVY = "0B132B"
INK = "17233C"
SLATE = "52627A"
MUTED = "7A879B"
PAPER = "F7F9FC"
WHITE = "FFFFFF"
TEAL = "0B8F8A"
TEAL_DARK = "086C68"
TEAL_LIGHT = "DDF3F1"
BLUE = "3A86FF"
BLUE_LIGHT = "E2EEFF"
ORANGE = "FF9F1C"
ORANGE_LIGHT = "FFF0D8"
CORAL = "EF476F"
CORAL_LIGHT = "FDE3EA"
GREEN = "2CB67D"
GREEN_LIGHT = "DDF5EA"
LAVENDER = "E9E7FF"
GRAY_LINE = "D8DFEA"
DARK_BG = "071426"

FONT_CN = "Microsoft YaHei"
FONT_EN = "Aptos"


def rgb(hex_color):
    from pptx.dml.color import RGBColor

    value = hex_color.strip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def ensure_assets():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    make_response_surface(ASSET_DIR / "position_response_surface.png")
    make_id_ood_scatter(ASSET_DIR / "position_id_ood_pareto.png")
    make_tradeoff_curve(ASSET_DIR / "small_data_tradeoff.png")


def make_response_surface(path):
    cmap = LinearSegmentedColormap.from_list(
        "tro_teal", ["#EDF7F6", "#A9DCD7", "#4CAEA8", "#0B8F8A", "#075B58"]
    )
    fig, ax = plt.subplots(figsize=(8.2, 5.2), dpi=220)
    fig.patch.set_facecolor("white")
    im = ax.imshow(COMBINED, cmap=cmap, vmin=0.18, vmax=0.52, aspect="auto")
    for i in range(3):
        for j in range(3):
            color = "white" if COMBINED[i, j] > 0.40 else "#17233C"
            ax.text(
                j,
                i,
                f"{COMBINED[i, j] * 100:.1f}%",
                ha="center",
                va="center",
                fontsize=17,
                fontweight="bold",
                color=color,
            )
    ax.set_xticks(range(3), ["3.6", "6.0", "8.4"], fontsize=12)
    ax.set_yticks(range(3), ["15", "25", "35"], fontsize=12)
    ax.set_xlabel("Approach radius (cm)", fontsize=13, fontweight="bold")
    ax.set_ylabel("Start radius (cm)", fontsize=13, fontweight="bold")
    ax.set_title(
        "Common-absolute deployment success", fontsize=17, fontweight="bold", pad=14
    )
    for spine in ax.spines.values():
        spine.set_visible(False)
    ax.set_xticks(np.arange(-0.5, 3, 1), minor=True)
    ax.set_yticks(np.arange(-0.5, 3, 1), minor=True)
    ax.grid(which="minor", color="white", linewidth=3)
    ax.tick_params(which="minor", bottom=False, left=False)
    cbar = fig.colorbar(im, ax=ax, fraction=0.044, pad=0.04)
    cbar.ax.tick_params(labelsize=10)
    cbar.outline.set_visible(False)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def make_id_ood_scatter(path):
    fig, ax = plt.subplots(figsize=(8.4, 5.4), dpi=220)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("#F7F9FC")
    approach_colors = {3.6: "#0B8F8A", 6.0: "#FF9F1C", 8.4: "#8A94A8"}
    offsets = {
        (15, 3.6): (-2, 9),
        (15, 6.0): (8, -12),
        (15, 8.4): (8, 8),
        (25, 3.6): (8, -13),
        (25, 6.0): (8, 8),
        (25, 8.4): (8, 8),
        (35, 3.6): (-35, -15),
        (35, 6.0): (8, 8),
        (35, 8.4): (-38, 8),
    }
    for i, start in enumerate(STARTS):
        for j, app in enumerate(APPROACHES):
            is_pareto = (start, app) in {(15, 3.6), (25, 6.0)}
            ax.scatter(
                ID_RATE[i, j],
                OOD_RATE[i, j],
                s=190 if is_pareto else 115,
                c=approach_colors[app],
                edgecolor="#EF476F" if is_pareto else "white",
                linewidth=3.0 if is_pareto else 1.4,
                zorder=4,
            )
            dx, dy = offsets[(start, app)]
            ax.annotate(
                f"{start}/{app:g}",
                (ID_RATE[i, j], OOD_RATE[i, j]),
                xytext=(dx, dy),
                textcoords="offset points",
                fontsize=10.5,
                color="#17233C",
                fontweight="bold" if is_pareto else "normal",
            )
    # Pareto frontier between the two non-dominated settings.
    ax.plot(
        [ID_RATE[1, 1], ID_RATE[0, 0]],
        [OOD_RATE[1, 1], OOD_RATE[0, 0]],
        color="#EF476F",
        linewidth=2.2,
        linestyle="--",
        alpha=0.85,
        zorder=2,
    )
    ax.set_xlim(0.12, 0.82)
    ax.set_ylim(0.11, 0.40)
    ax.set_xlabel("Condition-relative ID success", fontsize=13, fontweight="bold")
    ax.set_ylabel("Condition-relative OOD success", fontsize=13, fontweight="bold")
    ax.set_title(
        "Position ID–OOD trade-off (N=30)", fontsize=17, fontweight="bold", pad=14
    )
    ax.grid(True, color="#D8DFEA", linewidth=0.9)
    for spine in ax.spines.values():
        spine.set_visible(False)
    ax.tick_params(labelsize=10.5, colors="#52627A")
    ax.text(
        0.80,
        0.117,
        "→ learnability",
        ha="right",
        va="bottom",
        fontsize=10,
        color="#52627A",
    )
    ax.text(
        0.126,
        0.392,
        "↑ robustness",
        ha="left",
        va="top",
        fontsize=10,
        color="#52627A",
    )
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def make_tradeoff_curve(path):
    x = np.linspace(0, 1, 300)
    learnability = 0.92 - 0.77 * x**1.35
    coverage = 0.14 + 0.78 * (1 - np.exp(-3.2 * x))
    utility = 0.54 * learnability + 0.46 * coverage - 0.10 * (x - 0.53) ** 2
    best = int(np.argmax(utility))
    fig, ax = plt.subplots(figsize=(7.8, 4.6), dpi=220)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    ax.plot(x, learnability, color="#3A86FF", linewidth=3, label="Consistency / fit")
    ax.plot(x, coverage, color="#FF9F1C", linewidth=3, label="Coverage / robustness")
    ax.plot(x, utility, color="#0B8F8A", linewidth=4, label="Deployment utility")
    ax.fill_between(
        x[max(best - 34, 0) : min(best + 35, len(x))],
        0,
        utility[max(best - 34, 0) : min(best + 35, len(x))],
        color="#DDF3F1",
        alpha=0.75,
    )
    ax.scatter([x[best]], [utility[best]], s=130, c="#EF476F", zorder=5)
    ax.annotate(
        "sweet region",
        (x[best], utility[best]),
        xytext=(20, 20),
        textcoords="offset points",
        fontsize=11,
        fontweight="bold",
        color="#EF476F",
        arrowprops=dict(arrowstyle="->", color="#EF476F", lw=1.8),
    )
    ax.set_xlabel("Demonstration variance  →", fontsize=12.5, fontweight="bold")
    ax.set_ylabel("Relative utility", fontsize=12.5, fontweight="bold")
    ax.set_xticks([0, 1], ["too consistent", "too diverse"])
    ax.set_yticks([])
    ax.set_ylim(0, 1.02)
    ax.grid(axis="x", color="#D8DFEA", linewidth=0.8)
    for spine in ax.spines.values():
        spine.set_visible(False)
    ax.legend(loc="lower center", bbox_to_anchor=(0.5, -0.33), ncol=3, frameon=False)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def build_deck(output):
    # python-pptx is installed into a temporary target for this task.
    from pptx import Presentation
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def set_bg(slide, color=PAPER):
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(color)

    def add_text(
        slide,
        text,
        x,
        y,
        w,
        h,
        size=20,
        color=INK,
        bold=False,
        font=FONT_CN,
        align=PP_ALIGN.LEFT,
        valign=MSO_ANCHOR.TOP,
        margin=0.03,
        line_spacing=1.0,
    ):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = Inches(margin)
        tf.margin_right = Inches(margin)
        tf.margin_top = Inches(margin)
        tf.margin_bottom = Inches(margin)
        tf.vertical_anchor = valign
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = align
        p.line_spacing = line_spacing
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = rgb(color)
        return box

    def add_runs(
        slide,
        runs,
        x,
        y,
        w,
        h,
        size=20,
        color=INK,
        align=PP_ALIGN.LEFT,
        valign=MSO_ANCHOR.TOP,
    ):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = Inches(0.03)
        tf.margin_right = Inches(0.03)
        tf.margin_top = Inches(0.03)
        tf.margin_bottom = Inches(0.03)
        tf.vertical_anchor = valign
        p = tf.paragraphs[0]
        p.alignment = align
        for idx, spec in enumerate(runs):
            run = p.add_run() if idx else p.runs[0]
            run.text = spec["text"]
            run.font.name = spec.get("font", FONT_CN)
            run.font.size = Pt(spec.get("size", size))
            run.font.bold = spec.get("bold", False)
            run.font.color.rgb = rgb(spec.get("color", color))
        return box

    def add_rect(
        slide,
        x,
        y,
        w,
        h,
        fill=WHITE,
        line=GRAY_LINE,
        radius=True,
        line_width=1.0,
    ):
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(
            shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
        return shape

    def add_circle(slide, cx, cy, d, fill, line=None, line_width=1.0):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(cx), Inches(cy), Inches(d), Inches(d)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.line.color.rgb = rgb(line or fill)
        shape.line.width = Pt(line_width)
        return shape

    def add_line(slide, x1, y1, x2, y2, color=GRAY_LINE, width=1.5, dash=False):
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(x1),
            Inches(y1),
            Inches(x2),
            Inches(y2),
        )
        line.line.color.rgb = rgb(color)
        line.line.width = Pt(width)
        if dash:
            line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        return line

    def add_chevron(slide, x, y, w=0.35, h=0.5, fill=GRAY_LINE):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.line.fill.background()
        return shape

    def add_pill(
        slide,
        text,
        x,
        y,
        w,
        h=0.35,
        fill=TEAL_LIGHT,
        color=TEAL_DARK,
        size=11,
        bold=True,
    ):
        add_rect(slide, x, y, w, h, fill=fill, line=fill, radius=True)
        return add_text(
            slide,
            text,
            x,
            y + 0.01,
            w,
            h - 0.02,
            size=size,
            color=color,
            bold=bold,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
            margin=0,
        )

    def add_bullet_list(
        slide,
        items,
        x,
        y,
        w,
        h,
        size=18,
        color=INK,
        bullet_color=TEAL,
        gap=0.16,
    ):
        cursor = y
        per = (h - gap * max(len(items) - 1, 0)) / max(len(items), 1)
        for item in items:
            add_circle(slide, x, cursor + 0.11, 0.11, bullet_color)
            add_text(
                slide,
                item,
                x + 0.25,
                cursor,
                w - 0.25,
                per,
                size=size,
                color=color,
                line_spacing=1.08,
            )
            cursor += per + gap

    def slide_header(slide, section, title, page, subtitle=None, dark=False):
        if dark:
            add_pill(
                slide,
                section.upper(),
                0.72,
                0.48,
                1.75,
                fill="14365A",
                color="9EEAE4",
                size=10,
            )
            title_color = WHITE
            sub_color = "AFC3DA"
            line_color = "1F3652"
        else:
            add_pill(slide, section.upper(), 0.72, 0.48, 1.75, size=10)
            title_color = INK
            sub_color = SLATE
            line_color = GRAY_LINE
        add_text(
            slide,
            title,
            0.72,
            0.92,
            11.8,
            0.58,
            size=27,
            color=title_color,
            bold=True,
        )
        if subtitle:
            add_text(
                slide,
                subtitle,
                0.74,
                1.48,
                11.65,
                0.36,
                size=12.5,
                color=sub_color,
            )
        add_line(slide, 0.72, 7.14, 12.62, 7.14, color=line_color, width=0.8)
        add_text(
            slide,
            "T-RO · Guidance Generator",
            0.74,
            7.18,
            4.0,
            0.22,
            size=8.5,
            color=sub_color,
        )
        add_text(
            slide,
            page,
            11.82,
            7.18,
            0.75,
            0.22,
            size=8.5,
            color=sub_color,
            align=PP_ALIGN.RIGHT,
        )

    def add_source(slide, text):
        add_text(
            slide,
            text,
            7.1,
            6.84,
            5.45,
            0.20,
            size=7.5,
            color=MUTED,
            align=PP_ALIGN.RIGHT,
        )

    def add_metric_card(
        slide,
        x,
        y,
        w,
        h,
        metric,
        label,
        accent=TEAL,
        fill=WHITE,
        note=None,
    ):
        add_rect(slide, x, y, w, h, fill=fill, line=fill)
        add_rect(slide, x, y, 0.08, h, fill=accent, line=accent, radius=False)
        add_text(slide, metric, x + 0.25, y + 0.19, w - 0.4, 0.48, 25, accent, True)
        add_text(slide, label, x + 0.25, y + 0.72, w - 0.4, 0.38, 12.5, INK, True)
        if note:
            add_text(slide, note, x + 0.25, y + 1.10, w - 0.4, h - 1.22, 10.5, SLATE)

    # Slide 1 — title
    slide = prs.slides.add_slide(blank)
    set_bg(slide, DARK_BG)
    add_circle(slide, 9.72, -1.52, 5.05, "0E746F", line="0E746F")
    add_circle(slide, 10.65, 0.36, 2.75, "11456B", line="11456B")
    add_circle(slide, 9.53, 4.56, 2.05, "A35E12", line="A35E12")
    add_pill(
        slide,
        "WEEKLY RESEARCH UPDATE · 28 JUL 2026",
        0.78,
        0.68,
        3.55,
        fill="14365A",
        color="9EEAE4",
        size=10,
    )
    add_text(
        slide,
        "从 Spatial Evidence\n到 Guidance Generator",
        0.78,
        1.45,
        8.65,
        1.72,
        size=34,
        color=WHITE,
        bold=True,
        line_spacing=0.95,
    )
    add_text(
        slide,
        "30 条 demonstrations 下，寻找 DP learner 的\nlearnability–generalisation sweet spots",
        0.82,
        3.48,
        7.75,
        1.02,
        size=19,
        color="C8D7E8",
        line_spacing=1.1,
    )
    for idx, (txt, fill, color) in enumerate(
        [
            ("POSITION", "173B63", "BFD9F5"),
            ("N = 30", "14365A", "9EEAE4"),
            ("ID vs OOD", "5E3710", "FFD59A"),
        ]
    ):
        add_pill(slide, txt, 0.82 + idx * 1.58, 5.08, 1.34, fill=fill, color=color)
    add_text(
        slide,
        "核心观点：variance 不是越小或越大越好，而是需要按 trajectory stage 与 deployment objective 生成。",
        0.82,
        6.12,
        10.75,
        0.54,
        size=14,
        color="AFC3DA",
    )
    add_text(
        slide,
        "T-RO research direction",
        10.26,
        6.75,
        2.25,
        0.28,
        size=9,
        color="AFC3DA",
        align=PP_ALIGN.RIGHT,
    )

    # Slide 2 — executive summary
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    slide_header(slide, "EXECUTIVE", "一页总览：本周完成了什么？", "02 / 14")
    cards = [
        (
            "WHY",
            "Guidance Generator",
            "目标是自动输出 stage-conditioned variance，而不是人工调一个全局 tolerance。",
            TEAL,
            TEAL_LIGHT,
        ),
        (
            "TEST",
            "Position 3 × 3",
            "Start × Approach 全因子；每格 5 个独立 dataset-policy repeats，每份仅 30 条 demonstrations。",
            BLUE,
            BLUE_LIGHT,
        ),
        (
            "FIND",
            "两个 sweet spots",
            "15/3.6 偏 ID learnability；25/6.0 偏 OOD robustness；近抓取 8.4 cm 全面较弱。",
            ORANGE,
            ORANGE_LIGHT,
        ),
        (
            "SO WHAT",
            "Generator 必须多目标",
            "不应输出“唯一最优 variance”；应给出 Pareto guidance，再按任务风险与人的负担选择。",
            CORAL,
            CORAL_LIGHT,
        ),
    ]
    for idx, (tag, title, body, accent, fill) in enumerate(cards):
        x = 0.75 + idx * 3.05
        add_rect(slide, x, 2.02, 2.72, 3.55, fill=WHITE, line=GRAY_LINE)
        add_pill(slide, tag, x + 0.22, 2.24, 0.86, fill=fill, color=accent, size=9.5)
        add_text(slide, title, x + 0.22, 2.91, 2.30, 0.70, 18, INK, True)
        add_text(slide, body, x + 0.22, 3.77, 2.28, 1.30, 13.3, SLATE, False, line_spacing=1.12)
    add_rect(slide, 0.75, 5.88, 11.77, 0.72, fill=NAVY, line=NAVY)
    add_text(
        slide,
        "本周 evidence 的价值：证明“trajectory stage × variance × deployment objective”之间存在可学习结构。",
        1.03,
        6.05,
        11.20,
        0.36,
        size=16,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    # Slide 3 — north star
    slide = prs.slides.add_slide(blank)
    set_bg(slide, DARK_BG)
    slide_header(
        slide,
        "NORTH STAR",
        "总纲：我们真正要学习的不是一个半径",
        "03 / 14",
        subtitle="目标：给定任务、阶段、learner 与数据预算，预测“哪里需要一致，哪里可以多样”。",
        dark=True,
    )
    pipeline = [
        ("任务 / 阶段\ncontext", "14365A"),
        ("Compatibility\nmodel", "0E746F"),
        ("Variance\nschedule", "5E3710"),
        ("Human\nGuidance", "54315D"),
        ("30 demos\n→ DP", "25405C"),
        ("ID / OOD\ndeployment", "743148"),
    ]
    x_positions = [0.72, 2.78, 4.84, 6.90, 8.96, 11.02]
    for idx, ((label, fill), x) in enumerate(zip(pipeline, x_positions)):
        add_rect(slide, x, 2.18, 1.56, 1.18, fill=fill, line=fill)
        add_text(
            slide,
            label,
            x + 0.08,
            2.42,
            1.40,
            0.70,
            size=14,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
        if idx < len(pipeline) - 1:
            add_chevron(slide, x + 1.68, 2.54, fill="39526F")
    add_line(slide, 11.79, 3.58, 1.50, 3.58, color="39526F", width=1.3, dash=True)
    add_text(
        slide,
        "deployment results 形成闭环监督信号",
        4.78,
        3.68,
        3.85,
        0.30,
        size=11,
        color="9FB5CC",
        align=PP_ALIGN.CENTER,
    )
    add_rect(slide, 0.94, 4.42, 11.45, 1.46, fill="0D2138", line="1F3652")
    add_text(slide, "Generator 的输出", 1.22, 4.72, 2.05, 0.34, 13, "9EEAE4", True)
    add_text(
        slide,
        "σ₁(Start), σ₂(Approach), σ₃(Contact), …",
        3.05,
        4.57,
        5.0,
        0.55,
        23,
        WHITE,
        True,
        font=FONT_EN,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "不是 single global σ",
        9.47,
        4.72,
        2.35,
        0.34,
        13,
        "FFD59A",
        True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "选择原则：满足 learnability 与 robustness 门槛，同时尽量减少对人类教师的限制。",
        1.22,
        5.28,
        10.58,
        0.34,
        12.5,
        "AFC3DA",
        align=PP_ALIGN.CENTER,
    )

    # Slide 4 — small-data trade-off
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    slide_header(
        slide,
        "HYPOTHESIS",
        "为什么只有 30 条数据时，variance 必须被设计？",
        "04 / 14",
    )
    slide.shapes.add_picture(
        str(ASSET_DIR / "small_data_tradeoff.png"),
        Inches(0.62),
        Inches(1.84),
        width=Inches(7.25),
    )
    add_rect(slide, 8.19, 1.92, 4.40, 1.26, fill=BLUE_LIGHT, line=BLUE_LIGHT)
    add_text(slide, "太一致", 8.46, 2.17, 1.28, 0.34, 18, BLUE, True)
    add_text(
        slide,
        "容易拟合，但 deployment shift 下脆弱",
        9.64,
        2.11,
        2.58,
        0.52,
        13,
        INK,
    )
    add_rect(slide, 8.19, 3.42, 4.40, 1.26, fill=ORANGE_LIGHT, line=ORANGE_LIGHT)
    add_text(slide, "太多样", 8.46, 3.67, 1.28, 0.34, 18, ORANGE, True)
    add_text(
        slide,
        "固定预算被摊薄，DP 难以学到精确动作",
        9.64,
        3.61,
        2.58,
        0.52,
        13,
        INK,
    )
    add_rect(slide, 8.19, 4.92, 4.40, 1.26, fill=TEAL_LIGHT, line=TEAL_LIGHT)
    add_text(slide, "甜点区", 8.46, 5.17, 1.28, 0.34, 18, TEAL, True)
    add_text(
        slide,
        "关键阶段一致；非关键阶段保留必要 coverage",
        9.64,
        5.11,
        2.58,
        0.52,
        13,
        INK,
    )
    add_text(
        slide,
        "研究假设：sweet region 随 trajectory stage 与 deployment target 改变。",
        0.92,
        6.36,
        11.50,
        0.38,
        16,
        TEAL_DARK,
        True,
        align=PP_ALIGN.CENTER,
    )

    # Slide 5 — experiment design
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    slide_header(
        slide,
        "THIS WEEK",
        "本周最小验证：只拆 Position 的 Start 与 Approach",
        "05 / 14",
    )
    stages = [
        ("Start", TEAL),
        ("Approach", BLUE),
        ("Descent", SLATE),
        ("Grasp", ORANGE),
        ("Lift", CORAL),
    ]
    for idx, (label, color) in enumerate(stages):
        x = 0.82 + idx * 2.37
        add_circle(slide, x, 1.88, 0.72, color)
        add_text(
            slide,
            str(idx + 1),
            x,
            2.01,
            0.72,
            0.36,
            15,
            WHITE,
            True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
        add_text(
            slide,
            label,
            x - 0.25,
            2.72,
            1.22,
            0.30,
            13,
            color,
            True,
            align=PP_ALIGN.CENTER,
        )
        if idx < 4:
            add_line(slide, x + 0.72, 2.24, x + 2.12, 2.24, color=GRAY_LINE, width=2.0)
    add_pill(slide, "本周只干预", 0.88, 3.25, 1.15, fill=CORAL_LIGHT, color=CORAL)
    add_text(
        slide,
        "Start variance × Approach variance",
        2.15,
        3.23,
        4.55,
        0.40,
        18,
        INK,
        True,
        font=FONT_EN,
    )
    # 3x3 condition grid
    grid_x, grid_y, cell_w, cell_h = 0.86, 3.92, 1.58, 0.72
    for j, app in enumerate(APPROACHES):
        add_text(
            slide,
            f"App {app:g} cm",
            grid_x + 1.10 + j * cell_w,
            grid_y - 0.45,
            cell_w,
            0.32,
            11,
            SLATE,
            True,
            align=PP_ALIGN.CENTER,
        )
    for i, start in enumerate(STARTS):
        add_text(
            slide,
            f"Start {start} cm",
            grid_x,
            grid_y + i * cell_h + 0.21,
            1.05,
            0.30,
            11,
            SLATE,
            True,
            align=PP_ALIGN.RIGHT,
        )
        for j in range(3):
            fill = TEAL_LIGHT if (i, j) in {(0, 0), (1, 1)} else WHITE
            line = TEAL if (i, j) in {(0, 0), (1, 1)} else GRAY_LINE
            add_rect(
                slide,
                grid_x + 1.16 + j * cell_w,
                grid_y + i * cell_h,
                cell_w - 0.08,
                cell_h - 0.08,
                fill=fill,
                line=line,
                radius=False,
                line_width=1.5 if line == TEAL else 0.8,
            )
            add_text(
                slide,
                "5 policies\n× 30 demos",
                grid_x + 1.18 + j * cell_w,
                grid_y + i * cell_h + 0.08,
                cell_w - 0.12,
                cell_h - 0.20,
                10.5,
                INK,
                True,
                align=PP_ALIGN.CENTER,
                valign=MSO_ANCHOR.MIDDLE,
            )
    add_metric_card(slide, 7.62, 3.66, 2.20, 2.13, "45", "independent policies", BLUE)
    add_metric_card(slide, 10.08, 3.66, 2.20, 2.13, "22,500", "validated rollout rows", TEAL)
    add_text(
        slide,
        "固定 learner、N=30、early stopping、checkpoint freeze；所有 policy 共享 deployment states 与 RNG。",
        7.66,
        6.06,
        4.58,
        0.58,
        12.5,
        SLATE,
        align=PP_ALIGN.CENTER,
    )
    add_source(slide, "Source: frozen Stage 2c composition protocol & merge")

    # Slide 6 — deployment surfaces
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    slide_header(
        slide,
        "EVALUATION",
        "两套 deployment surface：不要把 ID 与 OOD 混成一个分数",
        "06 / 14",
    )
    add_rect(slide, 0.76, 1.90, 5.72, 4.48, fill=WHITE, line=GRAY_LINE)
    add_pill(slide, "PRIMARY", 1.06, 2.18, 1.06, fill=BLUE_LIGHT, color=BLUE)
    add_text(slide, "Common-absolute deployment", 1.06, 2.78, 4.85, 0.42, 21, INK, True)
    add_text(
        slide,
        "所有条件在同一组绝对 Start states 上比较",
        1.06,
        3.27,
        4.92,
        0.38,
        13.5,
        SLATE,
    )
    add_line(slide, 1.19, 4.02, 5.84, 4.02, color=GRAY_LINE, width=4)
    add_circle(slide, 1.25, 3.82, 0.40, BLUE)
    add_circle(slide, 3.17, 3.82, 0.40, BLUE)
    add_circle(slide, 5.31, 3.82, 0.40, BLUE)
    add_text(slide, "Inner", 1.07, 4.38, 0.80, 0.28, 12, BLUE, True, align=PP_ALIGN.CENTER)
    add_text(slide, "same states", 2.54, 4.38, 1.75, 0.28, 12, SLATE, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Outer", 5.09, 4.38, 0.84, 0.28, 12, BLUE, True, align=PP_ALIGN.CENTER)
    add_text(
        slide,
        "用于公平 condition ranking 与 interaction analysis",
        1.06,
        5.24,
        4.96,
        0.68,
        15,
        BLUE,
        True,
        align=PP_ALIGN.CENTER,
    )
    add_rect(slide, 6.86, 1.90, 5.72, 4.48, fill=WHITE, line=GRAY_LINE)
    add_pill(slide, "SECONDARY", 7.16, 2.18, 1.16, fill=ORANGE_LIGHT, color=ORANGE)
    add_text(slide, "Condition-relative ID / OOD", 7.16, 2.78, 4.85, 0.42, 21, INK, True)
    add_text(
        slide,
        "每个 Start family 按自身训练 support 定义边界",
        7.16,
        3.27,
        4.92,
        0.38,
        13.5,
        SLATE,
    )
    add_line(slide, 7.29, 4.02, 11.94, 4.02, color=GRAY_LINE, width=4)
    add_circle(slide, 7.35, 3.82, 0.40, TEAL)
    add_circle(slide, 9.27, 3.82, 0.40, ORANGE)
    add_circle(slide, 11.41, 3.82, 0.40, ORANGE)
    add_text(slide, "ID", 7.17, 4.38, 0.80, 0.28, 12, TEAL, True, align=PP_ALIGN.CENTER)
    add_text(slide, "support edge", 8.66, 4.38, 1.62, 0.28, 12, SLATE, True, align=PP_ALIGN.CENTER)
    add_text(slide, "OOD", 11.18, 4.38, 0.84, 0.28, 12, ORANGE, True, align=PP_ALIGN.CENTER)
    add_text(
        slide,
        "用于判断 learnability–robustness trade-off",
        7.16,
        5.24,
        4.96,
        0.68,
        15,
        ORANGE,
        True,
        align=PP_ALIGN.CENTER,
    )
    add_rect(slide, 1.56, 6.55, 10.24, 0.32, fill=CORAL_LIGHT, line=CORAL_LIGHT)
    add_text(
        slide,
        "注意：不同 Start family 的 condition-relative banks 不是同一绝对 deployment distribution。",
        1.72,
        6.53,
        9.92,
        0.36,
        11.5,
        CORAL,
        True,
        align=PP_ALIGN.CENTER,
    )

    # Slide 7 — response surface
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    slide_header(
        slide,
        "RESULT 1",
        "Position response surface 不是单调的",
        "07 / 14",
        subtitle="数字为 common-absolute combined deployment success。",
    )
    slide.shapes.add_picture(
        str(ASSET_DIR / "position_response_surface.png"),
        Inches(0.52),
        Inches(1.78),
        width=Inches(7.35),
    )
    add_metric_card(
        slide,
        8.07,
        1.92,
        4.30,
        1.29,
        "49.4%",
        "最高 combined：Start 15 / Approach 3.6",
        TEAL,
        fill=TEAL_LIGHT,
    )
    add_metric_card(
        slide,
        8.07,
        3.42,
        4.30,
        1.29,
        "49.0%",
        "近似并列：Start 25 / Approach 6.0",
        BLUE,
        fill=BLUE_LIGHT,
    )
    add_rect(slide, 8.07, 4.92, 4.30, 1.31, fill=CORAL_LIGHT, line=CORAL_LIGHT)
    add_text(slide, "Approach 8.4 cm", 8.36, 5.13, 3.70, 0.32, 17, CORAL, True)
    add_text(
        slide,
        "在 Start 15 / 25 / 35 三行均为最弱列",
        8.36,
        5.54,
        3.65,
        0.42,
        13,
        INK,
    )
    add_text(
        slide,
        "结论：不能沿一个“variance 越大越 robust”的单调假设设计 guidance。",
        7.95,
        6.47,
        4.54,
        0.42,
        13.5,
        TEAL_DARK,
        True,
        align=PP_ALIGN.CENTER,
    )
    add_source(slide, "Source: Stage 2c common-absolute 3×3 analysis")

    # Slide 8 — Pareto
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    slide_header(
        slide,
        "RESULT 2",
        "出现两个互补的 Pareto sweet spots",
        "08 / 14",
        subtitle="右上越好；红圈是九个设置中不被同时支配的解。",
    )
    slide.shapes.add_picture(
        str(ASSET_DIR / "position_id_ood_pareto.png"),
        Inches(0.52),
        Inches(1.74),
        width=Inches(7.45),
    )
    add_rect(slide, 8.09, 1.93, 4.22, 1.62, fill=TEAL_LIGHT, line=TEAL)
    add_pill(slide, "ID / LEARNABILITY", 8.36, 2.14, 1.56, fill=WHITE, color=TEAL, size=9)
    add_text(slide, "Start 15 / Approach 3.6", 8.36, 2.66, 3.55, 0.34, 17, TEAL_DARK, True)
    add_text(slide, "ID 75.8%  ·  OOD 26.9%  ·  Combined 49.4%", 8.36, 3.05, 3.56, 0.30, 11.5, INK)
    add_rect(slide, 8.09, 3.80, 4.22, 1.62, fill=ORANGE_LIGHT, line=ORANGE)
    add_pill(slide, "OOD / ROBUSTNESS", 8.36, 4.01, 1.52, fill=WHITE, color=ORANGE, size=9)
    add_text(slide, "Start 25 / Approach 6.0", 8.36, 4.53, 3.55, 0.34, 17, "A35E12", True)
    add_text(slide, "ID 47.8%  ·  OOD 36.0%  ·  Combined 49.0%", 8.36, 4.92, 3.56, 0.30, 11.5, INK)
    add_rect(slide, 8.09, 5.72, 4.22, 0.86, fill=WHITE, line=GRAY_LINE)
    add_text(
        slide,
        "Start 35 / Approach 6.0 虽然 ID≈OOD，\n但绝对性能较低；“gap 小”不等于“好”。",
        8.34,
        5.91,
        3.72,
        0.50,
        11.8,
        SLATE,
        align=PP_ALIGN.CENTER,
    )
    add_source(slide, "Source: Stage 2c condition-relative RMAX40 analysis")

    # Slide 9 — scientific conclusion
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    slide_header(
        slide,
        "CONCLUSION",
        "本周 spatial evidence 支持什么？不支持什么？",
        "09 / 14",
    )
    add_text(slide, "SUPPORTED", 0.86, 1.95, 5.55, 0.32, 12, GREEN, True, font=FONT_EN)
    add_rect(slide, 0.76, 2.38, 5.82, 3.89, fill=GREEN_LIGHT, line=GREEN_LIGHT)
    supported = [
        "Start × Approach 存在 interaction：同一个 variance 的作用依赖 trajectory stage 组合。",
        "近抓取阶段过宽（8.4 cm）在测试网格中持续不利。",
        "最佳 guidance 取决于目标：ID learnability 与 OOD robustness 对应不同 sweet spot。",
    ]
    add_bullet_list(slide, supported, 1.05, 2.72, 5.13, 2.98, size=15, bullet_color=GREEN)
    add_text(slide, "BOUNDARY", 6.99, 1.95, 5.42, 0.32, 12, CORAL, True, font=FONT_EN)
    add_rect(slide, 6.86, 2.38, 5.72, 3.89, fill=CORAL_LIGHT, line=CORAL_LIGHT)
    boundary = [
        "不是连续空间中的全局最优，也不是唯一精确半径。",
        "目前仅是一个 grasping task、Position、固定 DP learner 与 N=30。",
        "还没有证明跨任务迁移、真实人类可实现性或 human teaching transfer。",
    ]
    add_bullet_list(slide, boundary, 7.16, 2.72, 5.02, 2.98, size=15, bullet_color=CORAL)
    add_rect(slide, 1.63, 6.54, 10.10, 0.36, fill=NAVY, line=NAVY)
    add_text(
        slide,
        "当前最强 claim：固定 learner 与 30-demo budget 下，Position variance 的 deployment effect 是 stage-conditioned。",
        1.80,
        6.54,
        9.76,
        0.36,
        11.7,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    # Slide 10 — architecture
    slide = prs.slides.add_slide(blank)
    set_bg(slide, DARK_BG)
    slide_header(
        slide,
        "GENERATOR",
        "如何把 spatial evidence 变成 Guidance Generator？",
        "10 / 14",
        subtitle="第一版采用可解释的 forward compatibility model，而不是小数据黑盒 inverse generator。",
        dark=True,
    )
    cols = [
        ("① Context", "task geometry\nscripted stage\nlearner + N\nID/OOD target", "14365A"),
        ("② Forward model", "predict p(ID)\npredict p(OOD)\nuncertainty\ninteraction", "0E746F"),
        ("③ Pareto selector", "success constraints\nrisk preference\nhuman burden\nconfidence", "5E3710"),
        ("④ Guidance", "corridor width\norientation tol.\nspeed window\nphase schedule", "54315D"),
    ]
    for idx, (title, body, fill) in enumerate(cols):
        x = 0.72 + idx * 3.12
        add_rect(slide, x, 2.05, 2.58, 3.26, fill=fill, line=fill)
        add_text(slide, title, x + 0.23, 2.35, 2.12, 0.36, 16, WHITE, True, align=PP_ALIGN.CENTER)
        add_line(slide, x + 0.33, 2.92, x + 2.25, 2.92, color="5F7A96", width=0.8)
        add_text(
            slide,
            body,
            x + 0.28,
            3.22,
            2.02,
            1.65,
            14,
            "D3DFEC",
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
            line_spacing=1.18,
        )
        if idx < 3:
            add_chevron(slide, x + 2.72, 3.38, 0.31, 0.62, fill="39526F")
    add_rect(slide, 1.08, 5.74, 11.18, 0.86, fill="0D2138", line="1F3652")
    add_text(
        slide,
        "为什么 forward model？同样的 success 可能对应多个 variance schedules；先预测结果，再选满足约束且对人最宽松的方案。",
        1.37,
        5.93,
        10.60,
        0.47,
        13.5,
        "C8D7E8",
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    # Slide 11 — objective and data
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    slide_header(
        slide,
        "MODEL",
        "Generator 学什么数据？如何选择 guidance？",
        "11 / 14",
    )
    add_rect(slide, 0.75, 1.90, 6.16, 4.72, fill=WHITE, line=GRAY_LINE)
    add_text(slide, "Compatibility table", 1.04, 2.18, 5.54, 0.36, 18, INK, True)
    fields = [
        ("task features", TEAL_LIGHT, TEAL_DARK),
        ("stage features", BLUE_LIGHT, BLUE),
        ("variation type + magnitude", ORANGE_LIGHT, "A35E12"),
        ("learner + data budget", LAVENDER, "5548A5"),
        ("deployment distribution", CORAL_LIGHT, CORAL),
        ("successes / trials", GREEN_LIGHT, GREEN),
    ]
    for idx, (name, fill, color) in enumerate(fields):
        row, col = divmod(idx, 2)
        x = 1.04 + col * 2.83
        y = 2.85 + row * 0.83
        add_rect(slide, x, y, 2.55, 0.60, fill=fill, line=fill)
        add_text(
            slide,
            name,
            x + 0.09,
            y + 0.13,
            2.37,
            0.32,
            12.2,
            color,
            True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
    add_text(
        slide,
        "第一版模型：regularised binomial / additive model / shallow tree\n→ 小数据、可解释、可输出 uncertainty",
        1.04,
        5.58,
        5.55,
        0.68,
        13,
        SLATE,
        align=PP_ALIGN.CENTER,
    )
    add_rect(slide, 7.18, 1.90, 5.40, 4.72, fill=NAVY, line=NAVY)
    add_pill(slide, "MULTI-OBJECTIVE SELECTOR", 7.50, 2.18, 2.18, fill="14365A", color="9EEAE4", size=9)
    add_text(
        slide,
        "max  λ · pID + (1−λ) · pOOD\n− β · guidance burden",
        7.48,
        2.94,
        4.75,
        1.08,
        23,
        WHITE,
        True,
        font=FONT_EN,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_line(slide, 7.72, 4.31, 12.05, 4.31, color="39526F", width=1.0)
    add_text(slide, "subject to", 7.50, 4.60, 1.10, 0.30, 11, "AFC3DA", True, font=FONT_EN)
    constraints = [
        "pID ≥ learnability threshold",
        "pOOD ≥ robustness threshold",
        "uncertainty ≤ confidence limit",
    ]
    add_bullet_list(slide, constraints, 7.57, 4.96, 4.36, 1.10, size=12.5, color="D3DFEC", bullet_color=ORANGE, gap=0.05)
    add_text(
        slide,
        "更稳妥的产品输出是 Pareto set + 推荐理由，而不是伪精确的单点。",
        7.53,
        6.18,
        4.70,
        0.32,
        11.7,
        "FFD59A",
        True,
        align=PP_ALIGN.CENTER,
    )

    # Slide 12 — roadmap
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    slide_header(
        slide,
        "ROADMAP",
        "一步步做：从单任务 evidence 到可验证的 guidance",
        "12 / 14",
    )
    steps = [
        ("0", "Spatial evidence", "当前完成", TEAL, True),
        ("1", "Multi-task Position", "scripted phases", BLUE, False),
        ("2", "Add Rotation / Velocity", "逐轴扩展", ORANGE, False),
        ("3", "Compatibility model", "learn forward map", "5548A5", False),
        ("4", "Held-out Task A", "freeze → predict → validate", CORAL, False),
        ("5", "Human A → unseen B", "realise / retain / transfer", GREEN, False),
    ]
    y0 = 2.05
    for idx, (num, title, sub, color, done) in enumerate(steps):
        y = y0 + idx * 0.76
        add_circle(slide, 0.87, y, 0.48, color)
        add_text(
            slide,
            "✓" if done else num,
            0.87,
            y + 0.06,
            0.48,
            0.30,
            13,
            WHITE,
            True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
        if idx < len(steps) - 1:
            add_line(slide, 1.11, y + 0.48, 1.11, y + 0.78, color=GRAY_LINE, width=2)
        add_text(slide, title, 1.59, y + 0.02, 3.20, 0.31, 15.5, INK, True, font=FONT_EN if idx else FONT_CN)
        add_text(slide, sub, 4.48, y + 0.03, 2.10, 0.30, 12, SLATE)
    add_rect(slide, 7.04, 1.99, 5.47, 4.58, fill=WHITE, line=GRAY_LINE)
    add_text(slide, "每一步都要有独立的 stop / go gate", 7.38, 2.29, 4.80, 0.36, 18, INK, True)
    gates = [
        ("Robot compatibility", "固定 learner / N；ID+OOD deployment"),
        ("Cross-task transfer", "held-out A 前冻结预测"),
        ("Guidance realisation", "人是否真的产生目标 variance"),
        ("Human learning", "撤除 guidance 后 retention"),
        ("Teaching transfer", "完全未见 Task B、无 guidance"),
    ]
    for idx, (title, sub) in enumerate(gates):
        y = 2.98 + idx * 0.64
        add_circle(slide, 7.40, y + 0.04, 0.20, TEAL if idx < 2 else ORANGE)
        add_text(slide, title, 7.78, y, 1.78, 0.27, 12.5, INK, True, font=FONT_EN)
        add_text(slide, sub, 9.58, y, 2.44, 0.30, 11.2, SLATE)
    add_text(
        slide,
        "原则：一次只扩大一个 claim；不把 robot-side success 直接等同于 human teaching success。",
        7.39,
        6.09,
        4.72,
        0.35,
        11.5,
        CORAL,
        True,
        align=PP_ALIGN.CENTER,
    )

    # Slide 13 — next actions / advisor decisions
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    slide_header(
        slide,
        "NEXT",
        "下一阶段的最小可执行计划 + 明天需要老师拍板",
        "13 / 14",
    )
    add_text(slide, "建议下一阶段", 0.83, 1.91, 5.62, 0.34, 18, TEAL_DARK, True)
    actions = [
        ("1", "选择 2–3 个 source tasks", "优先已有 scripted phases 的 grasp / insertion 类任务"),
        ("2", "Position-only 紧凑验证", "围绕已发现 sweet regions，不做无边界半径搜索"),
        ("3", "冻结 common + ID/OOD banks", "每条件 N=30、matched seeds、独立 policies"),
        ("4", "建立 compatibility table", "先拟合低复杂度 forward model 与 uncertainty"),
    ]
    for idx, (num, title, sub) in enumerate(actions):
        y = 2.43 + idx * 0.95
        add_circle(slide, 0.87, y, 0.43, TEAL)
        add_text(slide, num, 0.87, y + 0.05, 0.43, 0.28, 12, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_text(slide, title, 1.53, y - 0.01, 2.55, 0.30, 14, INK, True)
        add_text(slide, sub, 4.00, y - 0.01, 2.35, 0.48, 11.5, SLATE)
    add_text(slide, "希望老师确认", 6.95, 1.91, 5.46, 0.34, 18, CORAL, True)
    decisions = [
        (
            "A",
            "Scope",
            "先明确为 DP-specific guidance，还是一开始就追求 learner-agnostic？",
        ),
        (
            "B",
            "Objective",
            "输出一个固定 λ 的 recommendation，还是输出 Pareto set 让任务方选择？",
        ),
        (
            "C",
            "Tasks",
            "哪 2–3 个 source tasks 最有论文说服力，同时工程上可控？",
        ),
    ]
    for idx, (tag, title, body) in enumerate(decisions):
        y = 2.43 + idx * 1.22
        add_rect(slide, 7.01, y, 5.27, 0.98, fill=CORAL_LIGHT if idx == 0 else WHITE, line=CORAL if idx == 0 else GRAY_LINE)
        add_pill(slide, tag, 7.25, y + 0.21, 0.52, h=0.42, fill=CORAL, color=WHITE, size=11)
        add_text(slide, title, 7.99, y + 0.18, 1.18, 0.30, 14, INK, True, font=FONT_EN)
        add_text(slide, body, 9.17, y + 0.13, 2.81, 0.60, 11.5, SLATE)
    add_rect(slide, 7.01, 6.10, 5.27, 0.54, fill=NAVY, line=NAVY)
    add_text(
        slide,
        "建议：先 DP-specific + Pareto output + scripted multi-task phases",
        7.20,
        6.20,
        4.88,
        0.30,
        11.5,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
    )

    # Slide 14 — take-home
    slide = prs.slides.add_slide(blank)
    set_bg(slide, DARK_BG)
    slide_header(slide, "TAKE-HOME", "今天希望带走的三句话", "14 / 14", dark=True)
    takeaways = [
        (
            "01",
            "Variance 是 stage-conditioned control variable",
            "不是需要被统一最小化或最大化的噪声。",
            TEAL,
        ),
        (
            "02",
            "30 条 demonstrations 也能暴露 Pareto structure",
            "15/3.6 偏 learnability；25/6.0 偏 robustness；8.4 cm near-grasp 明显不利。",
            ORANGE,
        ),
        (
            "03",
            "下一步不是继续调 grasping 半径",
            "而是跨任务学习 compatibility，冻结预测，再把它翻译成可实现的 human guidance。",
            CORAL,
        ),
    ]
    for idx, (num, title, body, accent) in enumerate(takeaways):
        y = 1.90 + idx * 1.47
        add_text(slide, num, 0.88, y, 0.70, 0.50, 25, accent, True, font=FONT_EN)
        add_line(slide, 1.72, y + 0.08, 1.72, y + 0.91, color="314A65", width=1.0)
        add_text(slide, title, 2.05, y - 0.01, 5.25, 0.38, 19, WHITE, True)
        add_text(slide, body, 2.05, y + 0.47, 8.68, 0.47, 13.5, "AFC3DA")
    add_rect(slide, 1.66, 6.46, 10.02, 0.43, fill="0D2138", line="1F3652")
    add_text(
        slide,
        "Discussion: 让 generator 输出一个 setting，还是输出可解释的 Pareto guidance menu？",
        1.91,
        6.50,
        9.52,
        0.32,
        13,
        "FFD59A",
        True,
        align=PP_ALIGN.CENTER,
    )

    # Appendix 1 — exact tables
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    slide_header(slide, "APPENDIX", "完整 3 × 3 数值表", "A1")
    datasets = [
        ("Common combined", COMBINED, BLUE_LIGHT, BLUE),
        ("Condition-relative ID", ID_RATE, TEAL_LIGHT, TEAL),
        ("Condition-relative OOD", OOD_RATE, ORANGE_LIGHT, ORANGE),
    ]
    for block, (title, data, highlight_fill, highlight_color) in enumerate(datasets):
        x0 = 0.72 + block * 4.18
        add_text(slide, title, x0, 1.92, 3.75, 0.34, 15, INK, True, font=FONT_EN, align=PP_ALIGN.CENTER)
        for j, app in enumerate(APPROACHES):
            add_text(slide, f"A{app:g}", x0 + 0.93 + j * 0.94, 2.45, 0.88, 0.26, 10, SLATE, True, align=PP_ALIGN.CENTER)
        for i, start in enumerate(STARTS):
            add_text(slide, f"S{start}", x0, 2.88 + i * 0.77, 0.74, 0.26, 10, SLATE, True, align=PP_ALIGN.RIGHT)
            for j in range(3):
                is_best = data[i, j] >= data.max() - 1e-9
                fill = highlight_fill if is_best else WHITE
                line = highlight_color if is_best else GRAY_LINE
                add_rect(slide, x0 + 0.91 + j * 0.94, 2.73 + i * 0.77, 0.86, 0.62, fill=fill, line=line, radius=False)
                add_text(slide, f"{data[i,j]*100:.1f}", x0 + 0.91 + j * 0.94, 2.89 + i * 0.77, 0.86, 0.26, 11.5, highlight_color if is_best else INK, is_best, align=PP_ALIGN.CENTER)
        add_text(slide, "% success", x0 + 0.94, 5.20, 2.75, 0.26, 9.5, MUTED, align=PP_ALIGN.CENTER)
    add_rect(slide, 0.94, 5.86, 11.48, 0.72, fill=WHITE, line=GRAY_LINE)
    add_text(
        slide,
        "Pareto front = Start15/App3.6 + Start25/App6.0.  Common-absolute 是主比较；condition-relative ID/OOD 是次级诊断。",
        1.17,
        6.05,
        11.04,
        0.34,
        12.5,
        INK,
        True,
        align=PP_ALIGN.CENTER,
    )
    add_source(slide, "Values: experiments/tro_stage2_position_composition/analysis/summary.md")

    # Appendix 2 — evidence audit
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    slide_header(slide, "APPENDIX", "Evidence audit：为什么这些结果可用于下一阶段设计？", "A2")
    metrics = [
        ("9", "Position conditions", "完整 Start × Approach 3×3"),
        ("5", "dataset-policy seeds", "独立 policy realisations"),
        ("30", "demos / dataset", "固定小数据预算"),
        ("45", "selected checkpoints", "rollout 前冻结 hash"),
        ("22,500", "formal rows analysed", "common + RMAX40"),
        ("52/52", "relevant tests", "strict merge + pairing"),
    ]
    for idx, (metric, label, note) in enumerate(metrics):
        row, col = divmod(idx, 3)
        x = 0.78 + col * 4.12
        y = 1.96 + row * 1.55
        add_metric_card(slide, x, y, 3.70, 1.28, metric, label, [TEAL, BLUE, ORANGE][col], fill=WHITE, note=note)
    add_rect(slide, 0.78, 5.34, 11.95, 1.17, fill=NAVY, line=NAVY)
    add_text(slide, "仍然不能越过的边界", 1.08, 5.57, 2.26, 0.30, 14, "FFD59A", True)
    add_text(
        slide,
        "single task · Position only · fixed DP learner · discrete candidate grid · simulation deployment",
        3.46,
        5.55,
        8.72,
        0.34,
        14,
        WHITE,
        True,
        font=FONT_EN,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "下一阶段的使命是扩展 external validity，而不是继续挖掘当前网格中的“更漂亮”数字。",
        1.08,
        6.03,
        11.26,
        0.30,
        11.7,
        "AFC3DA",
        align=PP_ALIGN.CENTER,
    )
    add_source(slide, "Frozen evidence: Stage 2 axial, RMAX40 rerollout, Stage 2c composition")

    props = prs.core_properties
    props.title = "从 Spatial Evidence 到 Guidance Generator"
    props.subject = "T-RO weekly research update: N=30 Position ID/OOD evidence"
    props.author = "polymetics_demt"
    props.keywords = "Guidance Generator, Diffusion Policy, Position, ID, OOD, N=30"
    props.comments = (
        "Generated from frozen Stage 2c results. "
        "See experiments/tro_stage2_position_composition/analysis/summary.md."
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return prs


def validate_deck(path, expected_slides=16):
    from zipfile import ZipFile

    from pptx import Presentation

    prs = Presentation(str(path))
    if len(prs.slides) != expected_slides:
        raise RuntimeError(f"Expected {expected_slides} slides, got {len(prs.slides)}")
    if path.stat().st_size < 100_000:
        raise RuntimeError(f"PPTX unexpectedly small: {path.stat().st_size} bytes")
    with ZipFile(path, "r") as archive:
        bad = archive.testzip()
        if bad is not None:
            raise RuntimeError(f"Corrupt PPTX member: {bad}")
        slide_xml = [
            name
            for name in archive.namelist()
            if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        ]
        if len(slide_xml) != expected_slides:
            raise RuntimeError(
                f"Expected {expected_slides} slide XML files, got {len(slide_xml)}"
            )
    titles = []
    for slide in prs.slides:
        texts = [
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        ]
        titles.append(texts[1] if len(texts) > 1 else texts[0] if texts else "")
    return titles


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--deps",
        type=Path,
        default=Path("/tmp/tro_ppt_deps"),
        help="Temporary python-pptx target directory.",
    )
    args = parser.parse_args()
    if args.deps.exists():
        sys.path.insert(0, str(args.deps))
    os.environ.setdefault("MPLCONFIGDIR", "/tmp/tro_ppt_mplconfig")
    ensure_assets()
    build_deck(args.output)
    titles = validate_deck(args.output)
    print(f"Created {args.output}")
    print(f"Slides: {len(titles)}")
    print(f"Bytes: {args.output.stat().st_size}")
    for idx, title in enumerate(titles, 1):
        print(f"{idx:02d}: {title[:80]}")


if __name__ == "__main__":
    main()
